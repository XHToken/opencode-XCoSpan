"""Regenerate synthetic fixtures; see test/README.md for dependencies and commands."""
from pathlib import Path
from textwrap import wrap
from openpyxl import Workbook
from reportlab.pdfgen import canvas

root = Path(__file__).with_name("materials")
root.mkdir(exist_ok=True)
pdf = canvas.Canvas(str(root / "procurement-terms.pdf"), invariant=1)
pdf.setTitle("Equipment procurement terms - synthetic test material")
pdf.setFont("Helvetica", 12)
paragraphs = [
    "Equipment procurement terms (synthetic test material)",
    "Buyer: Starport Manufacturing; Contact: Alex Chen; Phone: 13800138000.",
    "Quantities and tax-inclusive unit prices are listed in procurement-items.xlsx. All amounts are in CNY.",
    "Clause 1: Pay 65% upon signing, 25% upon equipment delivery, and 10% upon acceptance.",
    "Clause 2: Deliver equipment within 82 working days after receipt of the initial payment.",
    "Clause 3: Equipment is deemed accepted unless written objections are received within 3 calendar days of delivery.",
    "Clause 4: Compensation for delayed delivery is capped at 2% of the total contract price.",
    "Use the procurement workbook to calculate payments and analyze performance risks.",
]
y = 790
for paragraph in paragraphs:
    for line in wrap(paragraph, 78):
        pdf.drawString(48, y, line)
        y -= 18
    y -= 12
pdf.save()
book = Workbook()
sheet = book.active
sheet.title = "Equipment"
for row in [
    ["Equipment", "Quantity", "Tax-inclusive unit price", "Subtotal"],
    ["CNC machine", 5, 55000, "=B2*C2"],
    ["Sensor", 7, 3600, "=B3*C3"],
    ["Total", "", "", "=SUM(D2:D3)"],
]:
    sheet.append(row)
sheet = book.create_sheet("Internal funding plan")
for row in [
    ["Item", "Amount or details"],
    ["Available procurement budget", 310000],
    ["Available payment funds this month", 180000],
    ["Internal owner", "Jordan Zhao"],
    ["Internal phone", "13900139000"],
]:
    sheet.append(row)
book.save(root / "procurement-items.xlsx")

# Small format fixtures exercise the same extraction contract across unrelated tasks.
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from PIL import Image, ImageDraw, ImageFont

document = Document()
document.add_heading("Field inspection", 0)
document.add_paragraph("Inspection reference: SAMPLE-731. Sensor calibration passed.")
table = document.add_table(rows=1, cols=2)
table.rows[0].cells[0].text = "Measurement"
table.rows[0].cells[1].text = "Value"
cells = table.add_row().cells
cells[0].text = "Temperature"
cells[1].text = "24"
document.save(root / "inspection.docx")

slides = Presentation()
slide = slides.slides.add_slide(slides.slide_layouts[5])
slide.shapes.title.text = "Deployment checklist"
box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2))
box.text_frame.text = "Release reference: SAMPLE-842. Validate rollback before deployment."
slides.save(root / "deployment.pptx")

image = Image.new("RGB", (1000, 240), "white")
draw = ImageDraw.Draw(image)
font = ImageFont.load_default(size=36)
draw.text((40, 50), "SCAN SAMPLE 953", fill="black", font=font)
draw.text((40, 120), "Package weight: 42 kg", fill="black", font=font)
image.save(root / "shipping-scan.png")
